"""
Ingestion: turn the four provided university documents into two artefacts.

  1. A STRUCTURED layer  -> courses.csv / minors.csv / baskets.csv
     Facts that must never be paraphrased by an LLM: course codes, credits,
     prerequisites, which semester a course is offered in, basket credit caps.
     These drive the deterministic eligibility engine.

  2. An UNSTRUCTURED layer -> chunks.jsonl
     Regulation prose from the Student Handbook and the student SOP, chunked
     with citation metadata (source, page, clause) so every generated sentence
     can be traced back to a document.

This split is the core design decision of the system: numbers are computed,
prose is retrieved. See report Section 3.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, asdict
from typing import Iterable

import openpyxl
import pandas as pd
import pdfplumber

from advisor import config

# --------------------------------------------------------------------------
# Excel: semester spread  ->  course catalogue
# --------------------------------------------------------------------------

_SEM_HEADERS = ("Course Code", "Course", "Pre-Req", "L", "T", "P", "C")
_NULLS = {"", "nil", "none", "na", "n/a", "-", "--"}


def _clean(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def _num(v) -> float | None:
    s = _clean(v)
    if s.lower() in _NULLS:
        return None
    try:
        return float(s)
    except ValueError:
        m = re.search(r"\d+(?:\.\d+)?", s)
        return float(m.group()) if m else None


def _norm_code(raw: str) -> str:
    """'DATA 301' -> 'DATA301'; 'MATH401/ COMP401' -> 'MATH401/COMP401'."""
    s = _clean(raw).upper()
    s = re.sub(r"\s*/\s*", "/", s)
    return re.sub(r"([A-Z]{2,6})\s+(\d{3}[A-Z]?)", r"\1\2", s)


def _split_prereqs(raw: str) -> list[str]:
    """'MATH202' / 'MATH202, COMP201' / 'NIL' -> list of course codes."""
    s = _clean(raw)
    if s.lower() in _NULLS:
        return []
    parts = re.split(r"[,;/&]|\band\b|\+", s, flags=re.I)
    out = []
    for p in parts:
        p = _clean(p)
        if p and p.lower() not in _NULLS:
            out.append(_norm_code(p))
    return out


def _find_blocks(ws) -> list[tuple[int, str]]:
    """Locate each semester block. Returns [(first_col, semester_label), ...]."""
    header_row = None
    for r in range(1, min(6, ws.max_row) + 1):
        vals = [_clean(c.value) for c in ws[r]]
        if vals.count("Course Code") >= 1:
            header_row = r
            break
    if header_row is None:
        return []

    label_row = max(1, header_row - 1)
    blocks: list[tuple[int, str]] = []
    for c in range(1, ws.max_column + 1):
        if _clean(ws.cell(header_row, c).value) != "Course Code":
            continue
        # the S1/S2/... label sits in the row above, at this column or the next
        label = ""
        for probe in (c + 1, c, c + 2, c - 1):
            if 1 <= probe <= ws.max_column:
                cand = _clean(ws.cell(label_row, probe).value)
                if re.fullmatch(r"S\s*\d+", cand, flags=re.I) or "SUMMER" in cand.upper():
                    label = cand.upper().replace(" ", "")
                    break
        blocks.append((c, label))

    # fall back to positional S1..Sn if the sheet has no usable labels
    if not any(lbl for _, lbl in blocks):
        blocks = [(c, f"S{i + 1}") for i, (c, _) in enumerate(blocks)]
    return blocks


def parse_semester_spread(path=None) -> pd.DataFrame:
    """Every 'Sem Spread*' sheet -> one row per (batch, semester, course)."""
    path = path or config.SOURCES["structure"]
    wb = openpyxl.load_workbook(path, data_only=True)
    rows: list[dict] = []

    for sheet in wb.sheetnames:
        if not sheet.lower().startswith("sem"):
            continue
        ws = wb[sheet]
        batch = _batch_label(sheet)
        blocks = _find_blocks(ws)
        if not blocks:
            continue
        header_row = _header_row(ws)

        basket_code = basket_name = ""
        for r in range(header_row + 1, ws.max_row + 1):
            # forward-fill the basket label from columns A/B
            a, b = _clean(ws.cell(r, 1).value), _clean(ws.cell(r, 2).value)
            if re.fullmatch(r"B\d+", a, flags=re.I):
                basket_code, basket_name = a.upper(), b
            elif a.lower().startswith("total"):
                basket_code = basket_name = ""

            for col, sem in blocks:
                code = _clean(ws.cell(r, col).value)
                title = _clean(ws.cell(r, col + 1).value)
                if not code or code.lower() in _NULLS or not title:
                    continue
                if code.lower().startswith("total") or code == "Course Code":
                    continue
                prereq_raw = _clean(ws.cell(r, col + 2).value)
                rows.append(
                    {
                        "batch": batch,
                        "sheet": sheet,
                        "semester": sem or "",
                        "basket_code": basket_code,
                        "basket": basket_name,
                        "course_code": _norm_code(code),
                        "title": title,
                        "prereq_raw": prereq_raw or "NIL",
                        "prereqs": _split_prereqs(prereq_raw),
                        "L": _num(ws.cell(r, col + 3).value),
                        "T": _num(ws.cell(r, col + 4).value),
                        "P": _num(ws.cell(r, col + 5).value),
                        "credits": _num(ws.cell(r, col + 6).value),
                    }
                )

    df = pd.DataFrame(rows)
    if df.empty:
        return df
    df["sem_num"] = df["semester"].str.extract(r"S(\d+)").astype("Float64")
    df = df.drop_duplicates(subset=["batch", "semester", "course_code"]).reset_index(drop=True)
    return df


def _header_row(ws) -> int:
    for r in range(1, min(6, ws.max_row) + 1):
        if any(_clean(c.value) == "Course Code" for c in ws[r]):
            return r
    return 1


def _batch_label(sheet: str) -> str:
    m = re.search(r"(20\d{2})", sheet)
    year = m.group(1) if m else "unknown"
    return f"{year}-DS" if "DS" in sheet.upper() else year


# --------------------------------------------------------------------------
# Excel: basket credit requirements  (Struct_* sheets)
# --------------------------------------------------------------------------

_BASKET_ALIASES = {
    "university core": "University Core",
    "univ core ": "University Core",
    "program core ": "Program Core",
    "internship/  capstone project": "Internship/Capstone",
    "internship/ capstone project": "Internship/Capstone",
    "specialization tracks ": "Specialization Tracks",
    "univ core": "University Core",
    "foundation": "Foundation",
    "program core": "Program Core",
    "programme core": "Program Core",
    "program honors": "Program Honors",
    "program hons": "Program Honors",
    "specialization tracks": "Specialization Tracks",
    "specialisation tracks": "Specialization Tracks",
    "minor/open": "Minor/Open",
    "open/minor": "Minor/Open",
    "internship/capstone project": "Internship/Capstone",
    "internship/capstone": "Internship/Capstone",
    "total credits": "TOTAL",
    "total": "TOTAL",
}


def _baskets_from_spread(path) -> list[dict]:
    """The 'Sem Spread' sheets carry a 'Fixed Min' column next to each basket
    (B1..B7) and a 'Total' row. Used to fill gaps in the Struct_* sheets."""
    wb = openpyxl.load_workbook(path, data_only=True)
    rows: list[dict] = []
    for sheet in wb.sheetnames:
        if not sheet.lower().startswith("sem"):
            continue
        ws, batch = wb[sheet], _batch_label(sheet)
        for r in range(1, ws.max_row + 1):
            a, b = _clean(ws.cell(r, 1).value), _clean(ws.cell(r, 2).value)
            val = _num(ws.cell(r, 4).value)
            if val is None:
                continue
            if re.fullmatch(r"B\d+", a, flags=re.I) and b:
                name = _BASKET_ALIASES.get(b.lower().strip(), b.strip())
                rows.append({"batch": batch, "basket": name, "min_credits": val})
            elif a.lower().startswith("total"):
                rows.append({"batch": batch, "basket": "TOTAL", "min_credits": val})
    return rows


def parse_baskets(path=None) -> pd.DataFrame:
    """Struct_* sheets -> minimum credits required per basket, per batch."""
    path = path or config.SOURCES["structure"]
    wb = openpyxl.load_workbook(path, data_only=True)
    rows: list[dict] = []
    for sheet in wb.sheetnames:
        if not sheet.lower().startswith("struct"):
            continue
        ws = wb[sheet]
        batch = _batch_label(sheet)
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                key = _clean(ws.cell(r, c).value).lower().rstrip(":")
                if key in _BASKET_ALIASES:
                    for probe in (c + 1, c + 2):
                        val = _num(ws.cell(r, probe).value)
                        if val is not None:
                            rows.append(
                                {
                                    "batch": batch,
                                    "basket": _BASKET_ALIASES[key],
                                    "min_credits": val,
                                }
                            )
                            break
    rows += _baskets_from_spread(path)
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    # A sheet contains several "Total Credits" cells (per-track subtotals and the
    # programme grand total). Only the largest is the degree requirement.
    total = (df[df.basket == "TOTAL"].groupby("batch", as_index=False)["min_credits"].max())
    rest = (df[df.basket != "TOTAL"].drop_duplicates(subset=["batch", "basket"]))
    if not total.empty:
        total["basket"] = "TOTAL"
    return pd.concat([rest, total], ignore_index=True).reset_index(drop=True)


# --------------------------------------------------------------------------
# Excel: minors
# --------------------------------------------------------------------------

def parse_minors(path=None) -> pd.DataFrame:
    """Each sheet of the minors workbook is one minor programme."""
    path = path or config.SOURCES["minors"]
    wb = openpyxl.load_workbook(path, data_only=True)
    rows: list[dict] = []
    code_re = re.compile(r"^[A-Z]{2,5}\s?\d{3}[A-Z]?$")
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for r in range(1, ws.max_row + 1):
            vals = [_clean(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            code = next((v for v in vals if code_re.match(_norm_code(v))), "")
            # the title is the longest alphabetic cell on the row
            texts = [v for v in vals if len(v) > 6 and not code_re.match(v.upper())
                     and not v.replace(".", "").isdigit()]
            title = max(texts, key=len) if texts else ""
            credits = next((_num(v) for v in vals if _num(v) is not None and 0 < _num(v) <= 6), None)
            if title and title.lower() not in _NULLS and not title.lower().startswith(("total", "course", "sl")):
                rows.append(
                    {
                        "minor": sheet.strip(),
                        "course_code": _norm_code(code),
                        "title": title,
                        "credits": credits,
                    }
                )
    return pd.DataFrame(rows).drop_duplicates().reset_index(drop=True)


# --------------------------------------------------------------------------
# PDF -> citation-carrying chunks
# --------------------------------------------------------------------------

@dataclass
class Chunk:
    chunk_id: str
    source: str          # human-readable document name, shown to the user
    doc_key: str         # 'handbook' | 'sop' | 'catalogue'
    page: int | None
    section: str         # top-level heading, e.g. "7. Attendance Requirements"
    clause: str          # most specific clause id the chunk starts at, e.g. "7.2"
    text: str
    heading: str = ""    # contextual header prepended before embedding
    parent: str = ""     # top-level clause id, for small-to-big expansion
    batch: str = ""      # curriculum batch, for metadata filtering (catalogue only)
    clause_end: str = "" # last clause id packed into this chunk

    def citation(self) -> str:
        bits = [self.source]
        if self.clause:
            if self.doc_key == "catalogue":
                bits.append(self.clause)
            elif self.clause_end and self.clause_end != self.clause:
                bits.append(f"Clause {self.clause}–{self.clause_end}")
            else:
                bits.append(f"Clause {self.clause}")
        if self.page:
            bits.append(f"p.{self.page}")
        return " — ".join(bits)

    def embed_text(self) -> str:
        """What actually gets embedded: heading + body (contextual chunk header).
        A bare sub-clause like '7.3 ... sixty five percent' is ambiguous on its
        own; prefixed with 'Attendance Requirements' it retrieves correctly."""
        return f"{self.heading}\n{self.text}" if self.heading else self.text


_LIGATURE = re.compile(r"\(cid:\d+\)")
_CLAUSE_RE = re.compile(r"^\s*(\d{1,2}(?:\.\d{1,2}){0,3})\.?\s+(?=[A-Z(‘'\"a-z])")
_TOPLEVEL_TITLE = re.compile(r"^\s*(\d{1,2})\.?\s+([A-Z][A-Za-z ,&/()\-–]{3,70})$")
_BULLET_RE = re.compile(r"^[•●▪\-]\s*(.+)$")
_TOC_RE = re.compile(r"(…{2,}|\.{5,})")
_SENT_SPLIT = re.compile(r"(?<=[.;:])\s+(?=[A-Z(a-z]\)|[A-Z(])")


def _pdf_pages(path) -> Iterable[tuple[int, str]]:
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            text = _LIGATURE.sub(" ", text)
            text = re.sub(r"[ \t]+", " ", text)
            yield i, text.strip()


def _words(s: str) -> int:
    return len(s.split())


def _split_long(text: str, max_words: int, overlap_sents: int = 1) -> list[str]:
    """Sentence-boundary split with a one-sentence overlap, for blocks that
    exceed the token budget on their own. Never cuts mid-sentence."""
    sents = [x.strip() for x in _SENT_SPLIT.split(text) if x.strip()]
    out, cur = [], []
    for snt in sents:
        if cur and _words(" ".join(cur + [snt])) > max_words:
            out.append(" ".join(cur))
            cur = cur[-overlap_sents:] if overlap_sents else []
        cur.append(snt)
    if cur:
        out.append(" ".join(cur))
    return out


def chunk_pdf(path, source: str, doc_key: str) -> list[Chunk]:
    """
    Structure-aware chunking in three passes.

    1. SEGMENT  the text into atomic blocks, one per clause ("7.2 ...") or
                SOP bullet ("• Course Add/Drop ..."). Table-of-contents pages
                are dropped: their dot-leader lines match every query
                lexically and carry no rules.
    2. PACK     consecutive blocks of the SAME top-level clause together up to
                a word budget (~300 tokens, inside the embedder's 512 limit).
                A chunk never crosses a top-level clause boundary, so its
                citation is always right. Oversized blocks are split on
                sentence boundaries with a one-sentence overlap.
    3. HEADER   each chunk gets a contextual header -- document > section
                title > clause -- that is embedded with the body, so short
                sub-clauses keep the meaning of the section they belong to.
    """
    max_words = config.CHUNK_MAX_WORDS
    blocks: list[dict] = []        # {page, top, top_title, clause, lines}
    top, top_title = "", ""
    cur = None

    def open_block(page, clause, title=None):
        nonlocal cur
        if cur and cur["lines"]:
            blocks.append(cur)
        cur = {"page": page, "top": top, "top_title": title or top_title,
               "clause": clause, "lines": []}

    for page_no, text in _pdf_pages(path):
        if not text:
            continue
        toc_lines = sum(1 for l in text.split("\n") if _TOC_RE.search(l))
        if toc_lines >= 3:                      # a contents page
            continue
        for line in text.split("\n"):
            line = line.strip()
            if not line or re.fullmatch(r"\d{1,3}", line):
                continue
            tt = _TOPLEVEL_TITLE.match(line)
            m = _CLAUSE_RE.match(line)
            b = _BULLET_RE.match(line) if doc_key == "sop" else None
            if tt:
                top, top_title = tt.group(1), f"{tt.group(1)}. {tt.group(2).strip()}"
                open_block(page_no, top, top_title)
            elif m:
                cid = m.group(1)
                if "." not in cid:
                    top = cid
                    head = line[m.end():].split(".")[0][:70].strip()
                    top_title = f"{cid}. {head}" if head else cid
                open_block(page_no, cid)
            elif b:
                title = b.group(1).split(".")[0][:70].strip()
                top, top_title = title, title
                open_block(page_no, "", title)
            elif cur is None:
                open_block(page_no, "")
            cur["lines"].append(line)
    if cur and cur["lines"]:
        blocks.append(cur)

    # -- pack --
    chunks: list[Chunk] = []
    n = 0

    def emit(page, section, clause, parent, body, clause_end=""):
        nonlocal n
        body = body.strip()
        if _words(body) < 8:
            return
        n += 1
        heading = f"{source} > {section}" if section else source
        chunks.append(Chunk(f"{doc_key}-{n:04d}", source, doc_key, page, section,
                            clause, body, heading=heading, parent=parent,
                            clause_end=clause_end))

    pack, pack_meta, pack_ids = [], None, []

    def last_id():
        ids = [i for i in pack_ids if i]
        return ids[-1] if ids else ""

    for blk in blocks:
        body = " ".join(blk["lines"])
        same_parent = pack_meta and pack_meta["top"] == blk["top"]
        if pack and (not same_parent or _words(" ".join(pack) + " " + body) > max_words):
            emit(pack_meta["page"], pack_meta["top_title"], pack_meta["clause"],
                 pack_meta["top"], " ".join(pack), last_id())
            pack, pack_meta, pack_ids = [], None, []
        if _words(body) > max_words:
            for piece in _split_long(body, max_words):
                emit(blk["page"], blk["top_title"], blk["clause"], blk["top"], piece)
            continue
        if not pack:
            pack_meta = blk
        pack.append(body)
        pack_ids.append(blk["clause"])
    if pack:
        emit(pack_meta["page"], pack_meta["top_title"], pack_meta["clause"],
             pack_meta["top"], " ".join(pack), last_id())

    # -- dedupe exact repeats (running headers, repeated notices) --
    seen, out = set(), []
    for c in chunks:
        key = re.sub(r"\W+", "", c.text.lower())[:400]
        if key not in seen:
            seen.add(key)
            out.append(c)
    return out


def catalogue_chunks(courses: pd.DataFrame, baskets: pd.DataFrame,
                     minors: pd.DataFrame) -> list[Chunk]:
    """
    Render the structured tables as short natural-language chunks so the
    retriever can also answer prose questions about courses ("what are the
    prerequisites for Statistics with R?") without a SQL-style query.
    """
    chunks, n = [], 0
    for (batch, code), grp in courses.groupby(["batch", "course_code"]):
        r = grp.iloc[0]
        sems = ", ".join(sorted({s for s in grp["semester"] if s}))
        pre = ", ".join(r["prereqs"]) if r["prereqs"] else "none"
        n += 1
        text = (
            f"Course {code} '{r['title']}' is part of the {batch} B.Tech curriculum, "
            f"basket '{r['basket'] or 'unspecified'}'. Credits: {r['credits']}. "
            f"L-T-P: {r['L']}-{r['T']}-{r['P']}. Prerequisite(s): {pre}. "
            f"Offered in semester(s): {sems or 'not listed'}."
        )
        chunks.append(Chunk(f"cat-{n:04d}", "Programme Structure & Semester Spread (Sept 2026)",
                            "catalogue", None, f"Batch {batch}", code, text,
                            heading=f"Course catalogue > batch {batch} > {code} {r['title']}",
                            parent=code, batch=str(batch)))

    for batch, grp in baskets.groupby("batch"):
        n += 1
        body = "; ".join(f"{r.basket}: {r.min_credits:g} credits" for r in grp.itertuples())
        chunks.append(Chunk(f"cat-{n:04d}", "Programme Structure & Semester Spread (Sept 2026)",
                            "catalogue", None, f"Batch {batch}", "credit-requirements",
                            f"Minimum credit requirements per basket for batch {batch}: {body}. "
                            f"The B.Tech degree requires these credits for graduation.",
                            heading=f"Programme structure > batch {batch} > credit requirements",
                            parent="credits", batch=str(batch)))

    for minor, grp in minors.groupby("minor"):
        n += 1
        listed = "; ".join(f"{r.course_code or '(code n/a)'} {r.title}" for r in grp.itertuples())
        chunks.append(Chunk(f"cat-{n:04d}", "Minor Courses for B.Tech Students", "catalogue",
                            None, "Minors", minor,
                            f"The {minor} minor for B.Tech students includes: {listed}.",
                            heading=f"Minor programmes > {minor}", parent="minors"))
    n += 1
    names = ", ".join(sorted(minors.minor.unique()))
    chunks.append(Chunk(f"cat-{n:04d}", "Minor Courses for B.Tech Students", "catalogue",
                        None, "Minors", "minor-list",
                        f"Minors available to B.Tech students: {names}.",
                        heading="Minor programmes > list of all minors", parent="minors"))
    return chunks


# --------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------

def build_all(verbose: bool = True) -> dict:
    courses = parse_semester_spread()
    baskets = parse_baskets()
    minors = parse_minors()

    courses.to_csv(config.COURSES_CSV, index=False)
    baskets.to_csv(config.BASKETS_CSV, index=False)
    minors.to_csv(config.MINORS_CSV, index=False)

    chunks = chunk_pdf(config.SOURCES["handbook"], "Student Handbook (August 2026)", "handbook")
    chunks += chunk_pdf(config.SOURCES["sop"], "SOP for Students (17 Aug 2026)", "sop")
    chunks += catalogue_chunks(courses, baskets, minors)

    with open(config.CHUNKS_JSONL, "w") as fh:
        for c in chunks:
            fh.write(json.dumps(asdict(c)) + "\n")

    if verbose:
        print(f"courses  : {len(courses):5d} rows  ({courses['course_code'].nunique()} unique codes, "
              f"{courses['batch'].nunique()} batches)")
        print(f"baskets  : {len(baskets):5d} rows")
        print(f"minors   : {len(minors):5d} rows  ({minors['minor'].nunique()} minors)")
        print(f"chunks   : {len(chunks):5d}  "
              f"(handbook {sum(c.doc_key=='handbook' for c in chunks)}, "
              f"sop {sum(c.doc_key=='sop' for c in chunks)}, "
              f"catalogue {sum(c.doc_key=='catalogue' for c in chunks)})")
    return {"courses": courses, "baskets": baskets, "minors": minors, "chunks": chunks}


def load_chunks() -> list[Chunk]:
    return [Chunk(**json.loads(l)) for l in open(config.CHUNKS_JSONL)]


if __name__ == "__main__":
    build_all()
